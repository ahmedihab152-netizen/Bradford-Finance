import json
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")


def main() -> None:
    files = sys.argv[1:]
    output = []
    for path in files:
        item = {
            "file": path,
            "exists": os.path.exists(path),
            "size": os.path.getsize(path) if os.path.exists(path) else None,
            "sheets": [],
        }
        if item["exists"]:
            try:
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
                for sheet in workbook.worksheets:
                    samples = []
                    for row in sheet.iter_rows(
                        min_row=1, max_row=min(sheet.max_row, 25), values_only=True
                    ):
                        values = list(row)
                        if any(value not in (None, "") for value in values):
                            samples.append(values[:20])
                        if len(samples) >= 6:
                            break
                    item["sheets"].append(
                        {
                            "name": sheet.title,
                            "rows": sheet.max_row,
                            "cols": sheet.max_column,
                            "samples": samples,
                        }
                    )
                workbook.close()
            except Exception as error:
                item["error"] = repr(error)
        output.append(item)
    print(json.dumps(output, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
