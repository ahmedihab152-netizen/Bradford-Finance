import hashlib
import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from docx import Document


def clean(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def text(value):
    value = clean(value)
    if value is None:
        return ""
    return re.sub(r"\.0$", "", str(value)).strip()


def piasters(value):
    try:
        return int(round(float(value or 0) * 100))
    except (TypeError, ValueError):
        return 0


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sql_json(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str).replace("'", "''")


def normalize_grade(value):
    raw = text(value).upper().replace(" ", "").replace(".", "")
    if raw in {"PREK", "PRK", "PRE-K"}:
        return "Pre-K", 0
    if raw == "KG1":
        return "KG1", 1
    if raw == "KG2":
        return "KG2", 2
    match = re.search(r"(?:GR|GRADE|G)(\d{1,2})", raw)
    if match:
        number = int(match.group(1))
        return f"Grade {number}", number + 2
    return text(value), 999


def chunks(values, size):
    for start in range(0, len(values), size):
        yield values[start:start + size]


def main():
    output_dir = Path(sys.argv[1])
    students_path, analysis_path, uniform_path, review_path = map(Path, sys.argv[2:6])
    output_dir.mkdir(parents=True, exist_ok=True)

    book = openpyxl.load_workbook(students_path, read_only=True, data_only=True)
    sheet = book["الطلبة"]
    students = []
    accounts = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=7, max_row=1580, max_col=328, values_only=True), 7):
        values = list(row)
        code, name = text(values[6]), text(values[7])
        if not code or not name or not re.search(r"\d", code):
            continue
        grade, grade_order = normalize_grade(values[9])
        joined = text(values[10])
        status = "INACTIVE" if any(flag in joined for flag in ("ترك", "منقطع", "سحب", "منسحب")) else "ACTIVE"
        due, paid, remaining = map(piasters, (values[22], values[29], values[37]))
        line_total = 0
        for col in range(45, 325, 5):
            if col < len(values):
                line_total += max(0, piasters(values[col]))
        balance_difference = due - paid - remaining
        payment_difference = paid - line_total
        issues = []
        if abs(balance_difference) > 1:
            issues.append(f"SOURCE_BALANCE_DIFFERENCE:{balance_difference}")
        if abs(payment_difference) > 1:
            issues.append(f"SOURCE_PAYMENT_LINES_DIFFERENCE:{payment_difference}")
        students.append({
            "student_code": code,
            "student_name": name,
            "grade": grade,
            "source_grade": clean(values[9]),
            "grade_sort_order": grade_order,
            "joined_year": int(values[10]) if isinstance(values[10], (int, float)) else None,
            "status": status,
        })
        accounts.append({
            "student_code": code,
            "source_row": row_number,
            "original_tuition_piasters": piasters(values[11]),
            "discount_percent": float(values[13] or 0),
            "discount_type": clean(values[14]),
            "discount_piasters": piasters(values[15]),
            "net_tuition_piasters": piasters(values[16]),
            "activity_piasters": piasters(values[17]),
            "books_piasters": piasters(values[18]),
            "transport_piasters": piasters(values[19]),
            "other_piasters": piasters(values[20]),
            "prior_year_balance_piasters": piasters(values[21]),
            "total_due_piasters": due,
            "legacy_paid_piasters": paid,
            "legacy_remaining_piasters": remaining,
            "notes": clean(values[39]),
            "issues": issues,
        })
    book.close()

    uniform_book = openpyxl.load_workbook(uniform_path, read_only=True, data_only=True)
    uniform_sheet = uniform_book[uniform_book.sheetnames[0]]
    uniforms = []
    for row_number, row in enumerate(uniform_sheet.iter_rows(min_row=5, max_row=124, min_col=2, max_col=12, values_only=True), 5):
        values = list(row)
        if not text(values[2]) or piasters(values[4]) <= 0 or piasters(values[6]) <= 0:
            continue
        item_name = text(values[2])
        size_match = re.search(r"(?:مقاس|م)\s*([0-9]+(?:\s*[-–]\s*[0-9]+)?)", item_name)
        uniforms.append({
            "sku": f"UNI2627-{row_number:04d}",
            "item_name": item_name,
            "size": size_match.group(1).replace(" ", "") if size_match else None,
            "stock_quantity": int(values[9] if isinstance(values[9], (int, float)) else values[3] or 0),
            "average_cost_piasters": piasters(values[4]),
            "selling_price_piasters": piasters(values[6]),
            "source_row": row_number,
            "invoice_number": text(values[1]),
        })
    uniform_book.close()

    review = Document(review_path)
    review_rows = sum(max(0, len(table.rows) - 1) for table in review.tables[1:])
    source_files = [
        ("students_2627_official", students_path, "STUDENTS", len(students)),
        ("student_payment_analysis_2627", analysis_path, "SUMMARY", 1),
        ("uniform_2627_official", uniform_path, "INVENTORY", len(uniforms)),
        ("missing_students_review_2627", review_path, "REVIEW", review_rows),
    ]

    statements = []
    source_payload = [{"file_key": key, "file_name": path.name, "sha256": sha256(path), "category": category,
                       "size_bytes": path.stat().st_size, "row_count": count} for key, path, category, count in source_files]
    statements.append(f"""
begin;
create schema if not exists archive;
create table if not exists archive.bf_20261022_pre_official_2627_accounts as
  select * from public.student_accounts where academic_year='2026-2027';
update public.student_accounts set source_meta=coalesce(source_meta,'{{}}'::jsonb)||jsonb_build_object(
  'superseded',true,'superseded_at','2026-10-22','superseded_by','طلبة عام 26-27.xlsx'),updated_at=now()
where academic_year='2026-2027' and source_meta->>'superseded' is distinct from 'true';
update public.students set status='INACTIVE',updated_at=now()
where coalesce(student_code,'') not like 'BUS26-%';
update public.uniform_items set active=false;
insert into public.erp_source_files(file_key,file_name,sha256,category,academic_year,size_bytes,row_count,status,active,metadata)
select x.file_key,x.file_name,x.sha256,x.category,'2026-2027',x.size_bytes,x.row_count,'IMPORTED',true,
  jsonb_build_object('loaded_at','2026-10-22','authoritative',true)
from jsonb_to_recordset('{sql_json(source_payload)}'::jsonb)
as x(file_key text,file_name text,sha256 text,category text,size_bytes bigint,row_count integer)
on conflict(file_key) do update set file_name=excluded.file_name,sha256=excluded.sha256,category=excluded.category,
 academic_year=excluded.academic_year,size_bytes=excluded.size_bytes,row_count=excluded.row_count,status='IMPORTED',active=true,
 metadata=excluded.metadata,imported_at=now();
commit;
""")

    for part in chunks(students, 180):
        statements.append(f"""
insert into public.students(student_code,student_name,grade,status,joined_year,normalized_grade,grade_sort_order,
 grade_review_required,source_grade,program_review_required,updated_at)
select x.student_code,x.student_name,x.grade,x.status,x.joined_year,x.grade,x.grade_sort_order,
 x.grade_sort_order=999,x.source_grade,true,now()
from jsonb_to_recordset('{sql_json(part)}'::jsonb)
as x(student_code text,student_name text,grade text,status text,joined_year integer,grade_sort_order integer,source_grade text)
on conflict(student_code) do update set student_name=excluded.student_name,grade=excluded.grade,status=excluded.status,
 joined_year=excluded.joined_year,normalized_grade=excluded.normalized_grade,grade_sort_order=excluded.grade_sort_order,
 grade_review_required=excluded.grade_review_required,source_grade=excluded.source_grade,updated_at=now();
""")

    for part in chunks(accounts, 140):
        statements.append(f"""
insert into public.student_accounts(student_id,academic_year,tuition_piasters,discount_piasters,books_piasters,
 activity_piasters,transport_piasters,prior_year_balance_piasters,notes,original_tuition_piasters,
 net_tuition_piasters,other_piasters,total_due_piasters,legacy_paid_piasters,legacy_remaining_piasters,
 discount_percent,discount_type,source_external_id,source_meta,opening_paid_piasters,
 opening_remaining_piasters,opening_due_piasters,division_code,organization_id,updated_at)
select s.id,'2026-2027',x.net_tuition_piasters,x.discount_piasters,x.books_piasters,x.activity_piasters,
 x.transport_piasters,x.prior_year_balance_piasters,x.notes,x.original_tuition_piasters,x.net_tuition_piasters,
 x.other_piasters,x.total_due_piasters,x.legacy_paid_piasters,x.legacy_remaining_piasters,x.discount_percent,
 x.discount_type,'SRC2627-'||x.student_code,
 jsonb_build_object('source_file','طلبة عام 26-27.xlsx','source_row',x.source_row,'issues',x.issues,'superseded',false),
 x.legacy_paid_piasters,x.legacy_remaining_piasters,x.total_due_piasters,
 case when upper(coalesce(s.program,''))='IG' then 'IG' else 'AMERICAN' end,public.default_organization_id(),now()
from jsonb_to_recordset('{sql_json(part)}'::jsonb) as x(
 student_code text,source_row integer,original_tuition_piasters bigint,discount_percent numeric,discount_type text,
 discount_piasters bigint,net_tuition_piasters bigint,activity_piasters bigint,books_piasters bigint,
 transport_piasters bigint,other_piasters bigint,prior_year_balance_piasters bigint,total_due_piasters bigint,
 legacy_paid_piasters bigint,legacy_remaining_piasters bigint,notes text,issues jsonb)
join public.students s on s.student_code=x.student_code;
""")

    for part in chunks(uniforms, 120):
        statements.append(f"""
insert into public.uniform_items(sku,item_name,size,selling_price_piasters,stock_quantity,average_cost_piasters,active)
select x.sku,x.item_name,x.size,x.selling_price_piasters,x.stock_quantity,x.average_cost_piasters,true
from jsonb_to_recordset('{sql_json(part)}'::jsonb)
as x(sku text,item_name text,size text,selling_price_piasters bigint,stock_quantity integer,average_cost_piasters bigint,
 source_row integer,invoice_number text)
on conflict(sku) do update set item_name=excluded.item_name,size=excluded.size,
 selling_price_piasters=excluded.selling_price_piasters,stock_quantity=excluded.stock_quantity,
 average_cost_piasters=excluded.average_cost_piasters,active=true;
""")

    summary = {
        "students": len(students),
        "active_students": sum(row["status"] == "ACTIVE" for row in students),
        "inactive_students": sum(row["status"] != "ACTIVE" for row in students),
        "accounts": len(accounts),
        "accounts_with_issues": sum(bool(row["issues"]) for row in accounts),
        "total_due_piasters": sum(row["total_due_piasters"] for row in accounts),
        "total_paid_piasters": sum(row["legacy_paid_piasters"] for row in accounts),
        "total_remaining_piasters": sum(row["legacy_remaining_piasters"] for row in accounts),
        "uniform_items": len(uniforms),
        "uniform_stock": sum(row["stock_quantity"] for row in uniforms),
        "review_rows": review_rows,
    }
    statements.append(f"""
update public.erp_source_files set imported_count=row_count,imported_at=now()
where file_key in ('students_2627_official','student_payment_analysis_2627','uniform_2627_official','missing_students_review_2627');
insert into public.import_batches(batch_key,source_name,status,row_count,imported_count,duplicate_count,totals,notes,completed_at)
values('SOURCE-2627-20261022','ملفات الطلبة واليونيفورم 2026-2027','IMPORTED',{len(students)+len(uniforms)},
 {len(students)+len(uniforms)},0,'{sql_json(summary)}'::jsonb,
 'استيراد المصدر الرسمي؛ فروق المصدر محفوظة للمراجعة ولم تُصحح بصمت.',now())
on conflict(batch_key) do update set status='IMPORTED',row_count=excluded.row_count,imported_count=excluded.imported_count,
 totals=excluded.totals,notes=excluded.notes,completed_at=now();
""")

    for index, statement in enumerate(statements, 1):
        (output_dir / f"{index:03d}.sql").write_text(statement.strip() + "\n", encoding="utf-8")
    (output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"sql_files": len(statements), **summary}, ensure_ascii=False))


if __name__ == "__main__":
    main()
