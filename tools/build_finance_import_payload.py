import hashlib
import json
import math
import os
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")


def clean(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def number(value):
    if value in (None, "", "-", " - "):
        return 0.0
    try:
        return float(Decimal(str(value).replace(",", "").strip()))
    except (InvalidOperation, ValueError):
        return 0.0


def piasters(value):
    return int(round(number(value) * 100))


def text(value):
    value = clean(value)
    return "" if value is None else str(value)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(raw[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return None


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def grade(value):
    raw = text(value).upper().replace(" ", "").replace(".", "")
    if raw in {"PREK", "PRK", "PRE-K"}:
        return "Pre-K", 0
    if raw in {"KG1", "KG01"}:
        return "KG1", 1
    if raw in {"KG2", "KG02"}:
        return "KG2", 2
    match = re.search(r"(?:GR|GRADE|G)(\d{1,2})", raw)
    if match:
        value = int(match.group(1))
        return f"Grade {value}", value + 2
    return text(value) or None, 999


def row_hash(data):
    raw = json.dumps(data, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def add_raw(output, file_key, sheet, row_number, record_type, external_key, data):
    normalized = {str(key): clean(value) for key, value in data.items() if clean(value) is not None}
    output["raw_rows"].append(
        {
            "file_key": file_key,
            "sheet_name": sheet,
            "source_row": row_number,
            "record_type": record_type,
            "external_key": external_key,
            "normalized_data": normalized,
            "row_hash": row_hash(normalized),
        }
    )


def read_students(path, output):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["الطلبة"]
    for row_number, row in enumerate(sheet.iter_rows(min_row=7, max_col=164, values_only=True), 7):
        values = list(row)
        code = text(values[8])
        name = text(values[10])
        if not code or not name or not re.search(r"\d", code):
            continue
        code = re.sub(r"\.0$", "", code)
        normalized_grade, sort_order = grade(values[12])
        joined = clean(values[13])
        joined_text = text(joined)
        status = "INACTIVE" if any(word in joined_text for word in ("ترك", "سحب", "منسحب")) else "ACTIVE"
        student = {
            "student_code": code,
            "student_name": name,
            "grade": normalized_grade,
            "source_grade": clean(values[12]),
            "grade_sort_order": sort_order,
            "national_id": text(values[7]) or None,
            "joined_year": int(number(joined)) if number(joined) else None,
            "status": status,
            "source_row": row_number,
        }
        account = {
            "student_code": code,
            "academic_year": "2025-2026",
            "original_tuition_piasters": piasters(values[14]),
            "discount_percent": number(values[16]),
            "discount_type": clean(values[17]),
            "discount_piasters": piasters(values[18]),
            "net_tuition_piasters": piasters(values[19]),
            "activity_piasters": piasters(values[20]) + piasters(values[21]),
            "books_piasters": piasters(values[22]),
            "transport_piasters": piasters(values[23]),
            "other_piasters": piasters(values[24]),
            "prior_year_balance_piasters": piasters(values[25]),
            "total_due_piasters": piasters(values[26]),
            "legacy_paid_piasters": piasters(values[33]),
            "legacy_remaining_piasters": piasters(values[41]),
            "notes": clean(values[43]),
            "source_row": row_number,
        }
        output["students"].append(student)
        output["accounts"].append(account)
        add_raw(output, "students_2526", sheet.title, row_number, "STUDENT_ACCOUNT", code, {**student, **account})
        for col in range(50, min(len(values), 160), 5):
            amount = piasters(values[col - 1])
            if amount <= 0:
                continue
            transaction_date = parse_date(values[col])
            receipt = text(values[col + 1])
            bank_notice = text(values[col + 2])
            deposit_date = parse_date(values[col + 3])
            payment = {
                "student_code": code,
                "transaction_date": transaction_date or deposit_date or "2026-08-31",
                "amount_piasters": amount,
                "receipt_number": receipt or None,
                "bank_reference": bank_notice or None,
                "payment_method": "BANK_TRANSFER" if bank_notice else "CASH",
                "payment_category": "TUITION",
                "source_row": row_number,
                "source_col": col,
            }
            output["payments"].append(payment)
            add_raw(output, "students_2526", sheet.title, row_number, "STUDENT_PAYMENT", f"{code}:{col}", payment)

    app = book["ابلكيشن"]
    for row_number, row in enumerate(app.iter_rows(min_row=7, max_col=27, values_only=True), 7):
        values = list(row)
        code = re.sub(r"\.0$", "", text(values[2]))
        name = text(values[4])
        if not name:
            continue
        normalized_grade, _ = grade(values[6])
        application = {
            "source_external_id": f"APP2526-{row_number}",
            "status": text(values[7]) or "غير محدد",
            "student_code": code or None,
            "student_name": name,
            "grade": normalized_grade,
            "file_due_piasters": piasters(values[8]),
            "file_paid_piasters": piasters(values[9]),
            "total_paid_piasters": piasters(values[10]),
            "collector": clean(values[24]),
            "phone": clean(values[23]),
            "source_row": row_number,
        }
        output["applications"].append(application)
        add_raw(output, "students_2526", app.title, row_number, "APPLICATION", application["source_external_id"], application)
    book.close()


def read_school_treasury(path, output):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["خزينة المدرسة"]
    for row_number, row in enumerate(sheet.iter_rows(min_row=4, min_col=2, max_col=16, values_only=True), 4):
        values = list(row)
        transaction_date = parse_date(values[8])
        inflow = piasters(values[1])
        outflow = piasters(values[2])
        bank = piasters(values[3])
        description = clean(values[4])
        if not transaction_date and not description and not any((inflow, outflow, bank)):
            continue
        record = {
            "source_row": row_number,
            "transaction_date": transaction_date,
            "source_balance_piasters": piasters(values[0]),
            "inflow_piasters": inflow,
            "outflow_piasters": outflow,
            "bank_piasters": bank,
            "description": description,
            "document_text": clean(values[5]),
            "source_code": clean(values[6]),
            "payment_method_text": clean(values[7]),
            "expense_type": clean(values[10]),
            "company": clean(values[11]),
            "invoice_number": clean(values[12]),
            "withholding_piasters": piasters(values[13]),
            "reported_amount_piasters": piasters(values[14]),
        }
        output["treasury"].append(record)
        add_raw(output, "school_treasury", sheet.title, row_number, "TREASURY", str(row_number), record)
    book.close()


def read_amany_treasury(path, output):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["خزينة المدرسة"]
    empty_run = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=4, max_row=min(sheet.max_row, 25000), max_col=12, values_only=True), 4):
        values = list(row)
        if not any(value not in (None, "") for value in values):
            empty_run += 1
            if empty_run > 250:
                break
            continue
        empty_run = 0
        amounts = {
            "cash_piasters": piasters(values[2]),
            "instapay_piasters": piasters(values[3]),
            "vodafone_cash_piasters": piasters(values[4]),
            "visa_piasters": piasters(values[5]),
        }
        if not any(amounts.values()) and not clean(values[6]):
            continue
        record = {
            "source_row": row_number,
            "transaction_date": parse_date(values[11]),
            "source_balance_piasters": piasters(values[0]),
            **amounts,
            "description": clean(values[6]),
            "document_text": clean(values[10]),
        }
        output["amany_treasury"].append(record)
        add_raw(output, "amany_treasury", sheet.title, row_number, "AMANY_CASHBOX", str(row_number), record)
    book.close()


def find_header(sheet, wanted):
    for row_number in range(1, min(sheet.max_row, 25) + 1):
        row = [text(sheet.cell(row_number, col).value) for col in range(1, 20)]
        if all(any(token in value for value in row) for token in wanted):
            return row_number, row
    return None, []


def read_banks(path, file_key, academic_year, output):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in book.worksheets:
        header_row, headers = find_header(sheet, ("وارد", "منصرف"))
        if not header_row:
            continue
        incoming_col = next(i + 1 for i, value in enumerate(headers) if "وارد" in value)
        outgoing_col = next(i + 1 for i, value in enumerate(headers) if "منصرف" in value)
        description_col = next((i + 1 for i, value in enumerate(headers) if "البيان" in value), 6)
        receipt_col = next((i + 1 for i, value in enumerate(headers) if "الايصال" in value), description_col + 1)
        date_col = next((i + 1 for i, value in enumerate(headers) if "الايداع" in value or "التاريخ" in value), receipt_col + 1)
        student_col = next((i + 1 for i, value in enumerate(headers) if "كود الطالب" in value), date_col + 1)
        notes_col = next((i + 1 for i, value in enumerate(headers) if "ملاحظ" in value), student_col + 2)
        empty_run = 0
        max_needed = max(incoming_col, outgoing_col, description_col, receipt_col, date_col, student_col, notes_col)
        rows = sheet.iter_rows(min_row=header_row + 1, max_row=min(sheet.max_row, 15000), max_col=max_needed, values_only=True)
        for row_number, values in enumerate(rows, header_row + 1):
            incoming = piasters(values[incoming_col - 1])
            outgoing = piasters(values[outgoing_col - 1])
            description = clean(values[description_col - 1])
            transaction_date = parse_date(values[date_col - 1])
            if not incoming and not outgoing and not description and not transaction_date:
                empty_run += 1
                if empty_run > 250:
                    break
                continue
            empty_run = 0
            if not incoming and not outgoing:
                continue
            record = {
                "file_key": file_key,
                "academic_year": academic_year,
                "sheet_name": sheet.title,
                "source_row": row_number,
                "transaction_date": transaction_date,
                "credit_piasters": incoming,
                "debit_piasters": outgoing,
                "description": description,
                "bank_reference": clean(values[receipt_col - 1]),
                "student_code": clean(values[student_col - 1]),
                "notes": clean(values[notes_col - 1]),
            }
            output["bank_transactions"].append(record)
            add_raw(output, file_key, sheet.title, row_number, "BANK_TRANSACTION", f"{sheet.title}:{row_number}", record)
    book.close()


def read_bus(path, output):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["الباصات"]
    route = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, min_col=3, max_col=6, values_only=True), 1):
        seq, name, source_grade, fee_value = row
        seq = clean(seq)
        name = clean(name)
        source_grade = clean(source_grade)
        fee = piasters(fee_value)
        if name and text(source_grade) in {"المرحله", "المرحلة"}:
            route = text(name)
            continue
        if not name or not source_grade or not fee or not number(seq):
            continue
        normalized_grade, sort_order = grade(source_grade)
        record = {
            "source_row": row_number,
            "student_name": text(name),
            "grade": normalized_grade,
            "grade_sort_order": sort_order,
            "route_name": route or "مسار غير محدد",
            "fee_piasters": fee,
            "academic_year": "2026-2027",
        }
        output["bus"].append(record)
        add_raw(output, "bus_2627", sheet.title, row_number, "BUS_ENROLLMENT", str(row_number), record)
    book.close()


def read_summary(path, file_key, output):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in book.worksheets:
        empty_run = 0
        max_col = min(sheet.max_column, 80)
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15000), max_col=max_col, values_only=True), 1):
            values = [clean(value) for value in row]
            if not any(value is not None for value in values):
                empty_run += 1
                if empty_run > 250:
                    break
                continue
            empty_run = 0
            data = {str(index + 1): value for index, value in enumerate(values) if value is not None}
            add_raw(output, file_key, sheet.title, row_number, "SOURCE_ROW", str(row_number), data)
    book.close()


def main():
    files = {Path(path).name: path for path in sys.argv[2:]}
    output = {
        "source_files": [],
        "raw_rows": [],
        "students": [],
        "accounts": [],
        "payments": [],
        "applications": [],
        "treasury": [],
        "amany_treasury": [],
        "bank_transactions": [],
        "bus": [],
    }
    config = [
        ("طلبة عام 25-26 (1).xlsx", "students_2526", "STUDENTS", "2025-2026"),
        ("الخزينه - د-اماني.xlsx", "amany_treasury", "TREASURY", None),
        ("بيان بالطلبة والمصروفات.xlsx", "student_expense_statement", "SUMMARY", "2025-2026"),
        ("الخصومات.xlsx", "discount_summary", "SUMMARY", "2025-2026"),
        ("تحليل سداد مصروفات الطلبة 2025-2026 (1).xlsx", "payment_analysis_copy", "SUMMARY", "2025-2026"),
        ("انصبة 2025-2026.xlsx", "activity_shares", "SUMMARY", "2025-2026"),
        ("تحليل سداد مصروفات الطلبة 2025-2026.xlsx", "payment_analysis", "SUMMARY", "2025-2026"),
        ("عدد الطلبة حتي 31-8-2026.xlsx", "student_count_aug", "SUMMARY", "2025-2026"),
        ("الايراد.xlsx", "revenue_summary", "SUMMARY", "2025-2026"),
        ("الخزينه (1).xlsx", "school_treasury", "TREASURY", None),
        ("حركة البنك 2025-2026.xlsx", "bank_2526", "BANK", "2025-2026"),
        ("حركة البنك 2027-2026.xlsx", "bank_2627", "BANK", "2026-2027"),
        ("عدد الطلبة حتي 25-9-2026.xlsx", "student_count_sep", "SUMMARY", "2026-2027"),
        ("باص عام 2026-2027.xlsx", "bus_2627", "TRANSPORT", "2026-2027"),
    ]
    for file_name, file_key, category, academic_year in config:
        path = files[file_name]
        output["source_files"].append(
            {
                "file_key": file_key,
                "file_name": file_name,
                "sha256": sha256(path),
                "category": category,
                "academic_year": academic_year,
                "size_bytes": os.path.getsize(path),
            }
        )
    read_students(files["طلبة عام 25-26 (1).xlsx"], output)
    read_school_treasury(files["الخزينه (1).xlsx"], output)
    read_amany_treasury(files["الخزينه - د-اماني.xlsx"], output)
    read_banks(files["حركة البنك 2025-2026.xlsx"], "bank_2526", "2025-2026", output)
    read_banks(files["حركة البنك 2027-2026.xlsx"], "bank_2627", "2026-2027", output)
    read_bus(files["باص عام 2026-2027.xlsx"], output)
    for file_name, file_key, category, _ in config:
        if category == "SUMMARY":
            read_summary(files[file_name], file_key, output)
    duplicate_hashes = {}
    for source in output["source_files"]:
        if source["sha256"] in duplicate_hashes:
            source["duplicate_of"] = duplicate_hashes[source["sha256"]]
        else:
            duplicate_hashes[source["sha256"]] = source["file_key"]
            source["duplicate_of"] = None
    for source in output["source_files"]:
        source["row_count"] = sum(1 for row in output["raw_rows"] if row["file_key"] == source["file_key"])
    path = Path(sys.argv[1])
    path.write_text(json.dumps(output, ensure_ascii=False, default=str), encoding="utf-8")
    summary = {key: len(value) for key, value in output.items() if isinstance(value, list)}
    summary["output"] = str(path)
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
